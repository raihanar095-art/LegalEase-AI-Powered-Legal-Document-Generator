import streamlit as st
import requests
import re
import os
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

# Page Configuration
st.set_page_config(page_title="LegalEase", layout="centered")

# Utility Functions
def sanitize_text(text: str) -> str:
    text = re.sub(f'[\u201c\u201d]', '"', text)
    text = re.sub(f'[\u2018\u2019]', "'", text)
    return text

def format_html_preview(text: str) -> str:
    html = text.replace("\n", "<br>")
    return f"<div style='background-color:#1e1e1e; color:#ffffff; padding:20px; border-radius:10px; font-family:sans-serif;'>{html}</div>"

def format_docx(text: str, doc_type: str, terms_text: str = "") -> BytesIO:
    doc = Document()
    
    logo_path = os.path.join("Image", "Logo.png")
    if os.path.exists(logo_path):
        doc.add_image(logo_path, width=Inches(1.5))
        
    title = doc.add_heading(doc_type, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    paragraphs = text.split("\n")
    for para in paragraphs:
        if para.strip():
            doc.add_paragraph(para)
            
    if terms_text:
        doc.add_heading("Terms & Conditions Summary", level=2)
        terms_list = [t.strip() for t in terms_text.split(";") if t.strip()]
        if terms_list:
            table = doc.add_table(rows=1, cols=2)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'No.'
            hdr_cells[1].text = 'Term Clause'
            for idx, term in enumerate(terms_list, 1):
                row_cells = table.add_row().cells
                row_cells[0].text = str(idx)
                row_cells[1].text = term

    section = doc.sections[0]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.text = "LegalEase Inc. | contact@legalease.com | All Rights Reserved."
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

class PDFGenerator(FPDF):
    def header(self):
        logo_path = os.path.join("Image", "Logo.png")
        if os.path.exists(logo_path):
            self.image(logo_path, 10, 8, 33)
            self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.cell(0, 10, "LegalEase Inc. | contact@legalease.com | All Rights Reserved.", 0, 0, "C")

def format_pdf(text: str, doc_type: str) -> BytesIO:
    pdf = PDFGenerator()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, doc_type, ln=True, align="C")
    pdf.ln(5)
    
    pdf.set_font("Arial", size=10)
    clean_text = text.encode('latin-1', 'replace').decode('latin-1')
    for line in clean_text.split("\n"):
        pdf.multi_cell(0, 6, line)
        
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

# UI Layout
logo_path = os.path.join("Image", "Logo.png")
col1, col2, col3 = st.columns([1, 2, 1])
with col2:
    if os.path.exists(logo_path):
        st.image(logo_path, use_container_width=True)

st.markdown("<h2 style='text-align: center;'>AI Legal Document Generator</h2>", unsafe_allow_html=True)

# Form Inputs
document_type = st.text_input("Document Type (Ex: Agreement, Contract, NDA)")
parties = st.text_area("Parties Involved")
terms = st.text_area("Terms & Conditions (Use semicolons for bullet points)")
dates = st.text_input("Effective Date")

if "generated_text" not in st.session_state:
    st.session_state.generated_text = ""
if "show_edit" not in st.session_state:
    st.session_state.show_edit = False

if st.button("Generate Document"):
    if document_type and parties and terms and dates:
        try:
            response = requests.post("http://localhost:8000/generate", json={
                "document_type": document_type,
                "parties": parties,
                "terms": terms,
                "dates": dates
            })
            if response.status_code == 200:
                raw_text = response.json().get("document", "")
                st.session_state.generated_text = sanitize_text(raw_text)
                st.success("Document Generated Successfully!")
            else:
                st.error("Failed to generate document from server.")
        except Exception as e:
            st.error(f"Backend connection error: {e}")
    else:
        st.warning("Please fill in all fields.")

if st.session_state.generated_text:
    st.markdown(format_html_preview(st.session_state.generated_text), unsafe_allow_html=True)
    
    if st.button("Click to Edit Document"):
        st.session_state.show_edit = not st.session_state.show_edit

    if st.session_state.show_edit:
        edited_text = st.text_area("Edit Document Below:", st.session_state.generated_text, height=300)
        st.session_state.generated_text = edited_text

    col_txt, col_docx, col_pdf = st.columns(3)
    
    with col_txt:
        st.download_button(
            label="Download as .TXT",
            data=st.session_state.generated_text,
            file_name=f"{document_type.replace(' ', '_').lower()}.txt",
            mime="text/plain"
        )
    
    with col_docx:
        docx_data = format_docx(st.session_state.generated_text, document_type, terms)
        st.download_button(
            label="Download as .DOCX",
            data=docx_data,
            file_name=f"{document_type.replace(' ', '_').lower()}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
    with col_pdf:
        pdf_data = format_pdf(st.session_state.generated_text, document_type)
        st.download_button(
            label="Download as .PDF",
            data=pdf_data,
            file_name=f"{document_type.replace(' ', '_').lower()}.pdf",
            mime="application/pdf"
        )